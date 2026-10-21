"""Compose actual illustrated scenes, without classroom boards or speech bubbles."""
from __future__ import annotations

import hashlib
import json
import os
import tempfile
import wave
from dataclasses import replace
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from zhijiang.course_styles import effective_template
from zhijiang.deck_renderer import Page, PAGE_H, PAGE_W
from zhijiang.drama import DramaStyleDraft, source_sentences, validate_shots
from zhijiang.page_media import media_fingerprint, narrated_pages, overlay_generated_window
from zhijiang.presentation import PresentationError


def checked_asset(root, item, key):
    path = (root/item[key]).resolve()
    if not path.is_relative_to(root.resolve()) or not path.is_file():
        raise PresentationError('漫剧镜头素材缺失或路径无效。')
    digest = item.get('sha256', {}).get(item[key])
    if not digest or hashlib.sha256(path.read_bytes()).hexdigest() != digest:
        raise PresentationError('漫剧镜头素材字节与生成记录不一致。')
    return path


def cue_time(offset, cues, source_length, duration):
    """Use measured utterance timings, interpolate only inside that utterance."""
    cursor = 0
    for cue in cues:
        end = cursor+len(cue['text'])
        if cursor <= offset <= end:
            return cue['start']+(cue['end']-cue['start'])*(offset-cursor)/max(1, end-cursor)
        cursor = end
    return duration*offset/max(1, source_length)


def render_drama_deck(lesson, output, *, audio_files=None):
    output = Path(output)
    root = output.parent
    root.mkdir(parents=True, exist_ok=True)
    previews, media = root/'presentation-preview', root/'presentation-media'
    previews.mkdir(exist_ok=True)
    media.mkdir(exist_ok=True)
    story = lesson.deck_plan['comic_story']
    style = DramaStyleDraft.model_validate(lesson.deck_plan['course_style']['design'])
    if story.get('design') != style.model_dump(mode='json'):
        raise PresentationError('漫剧角色与场景设定已变化，需要重新规划。')
    # Plain cinematic canvas. Fixed PPT theme ornaments must not cover the scene.
    template = replace(effective_template(lesson), id='scene_drama', decoration='none',
                       background='101623', ink='FFFFFF', muted='CFD8E3')
    if audio_files is None:
        audio_files = [root/f'audio-{i}.wav' for i in range(len(lesson.segments))]
    if len(audio_files) != len(lesson.segments):
        raise PresentationError('漫剧旁白数量与原文片段不一致。')
    speech_path = root/'source-speech.json'
    speech = json.loads(speech_path.read_text(encoding='utf8')) if speech_path.is_file() else []
    speech = {str(v['segment_id']): v for v in speech}
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(PAGE_W), Inches(PAGE_H)
    deck.core_properties.title = lesson.title
    deck.core_properties.subject = 'Lively Class——AI教育Agent · 场景漫剧 · 教材原文旁白'
    pages, segments = [], []
    for index, (segment, audio) in enumerate(zip(lesson.segments, audio_files), 1):
        entry = story['segments'].get(str(index), {})
        if entry.get('narration_sha256') != hashlib.sha256(segment.narration.encode()).hexdigest():
            raise PresentationError('原文旁白已变化，需要重新规划漫剧。')
        shots = entry.get('shots', [])
        validate_shots(shots, range(1, len(source_sentences(segment.narration))+1), style)
        if ''.join(s['text'] for s in shots) != segment.narration:
            raise PresentationError('漫剧旁白出现缺失、改写或重复。')
        # A missing shot never silently becomes a teacher/student illustration.
        raw_assets = lesson.deck_plan.get('generated_media', {}).get(str(index), [])
        assets = {a.get('shot_id'): a for a in raw_assets}
        if len(assets) != len(raw_assets) or set(assets) != {s['id'] for s in shots}:
            raise PresentationError('漫剧需要每个镜头对应的真实绘制素材，请完成绘图或视频生成。')
        if not Path(audio).is_file():
            raise PresentationError('漫剧缺少原文旁白音频。')
        with wave.open(str(audio)) as w:
            duration = w.getnframes()/w.getframerate()
        timing = speech.get(str(index), {})
        cues = timing.get('cues', [])
        if cues and (''.join(c['text'] for c in cues) != segment.narration or
                     abs(timing['audio_duration']-duration) > .05):
            raise PresentationError('漫剧分镜时序与原文配音记录不一致。')
        states, numbers, overlays, fractions, records = [], [], [], [], []
        cursor = 0
        for shot in shots:
            if shot['source_start'] != cursor or shot['source_end'] != cursor+len(shot['text']):
                raise PresentationError('漫剧镜头的原文字符区间不连续。')
            if shot['text'] != segment.narration[shot['source_start']:shot['source_end']]:
                raise PresentationError('漫剧镜头文字与来源区间不一致。')
            asset = assets[shot['id']]
            poster = checked_asset(root, asset, 'file')
            page = Page(deck, template, previews, len(deck.slides)+1)
            # Keep the complete provider composition; never crop away characters.
            page.picture(poster, (0, 0, PAGE_W, 7.08))
            page.text(segment.title, (.24, 7.17, 9.3, .25), size=13, min_size=13)
            page.text(f'AI情境演绎 · 教材第{segment.evidence.page}页', (9.6, 7.17, 3.45, .25),
                      size=11, min_size=11, color=template.muted, align='right')
            page.slide.notes_slide.notes_text_frame.text = (
                '教材原文旁白（完整片段）：\n'+segment.narration+'\n\n本镜头原句：\n'+shot['text']+
                '\n\n分镜：\n'+json.dumps(shot, ensure_ascii=False, indent=2)+
                '\nAI绘制的情境演绎，非教材原图。')
            states.append(page.save_preview())
            numbers.append(page.number)
            fractions.append(cue_time(cursor, cues, len(segment.narration), duration)/duration)
            cursor = shot['source_end']
            if asset.get('video'):
                overlays.append((len(states)-1, checked_asset(root, asset, 'video')))
            records.append({**shot, 'page': page.number, 'asset': asset['file'],
                'actual_medium': 'provider_video' if asset.get('video') else 'generated_still'})
            pages.append({'page': page.number, 'kind': 'drama_shot', 'segment_id': index, 'shot_id': shot['id']})
        clip = media/f'segment-{index:03d}.mp4'
        narrated_pages(states, audio, clip, fractions=fractions)
        for n, source in overlays:
            temporary = media/f'drama-overlay-{index}-{n}.mp4'
            end = fractions[n+1] if n+1 < len(fractions) else 1
            overlay_generated_window(clip, source, (0, 0, 1280, 680), temporary,
                                     fractions[n]*duration, end*duration, duration)
            os.replace(temporary, clip)
        segments.append({'segment_id': index, 'pages': numbers, 'clip': clip.relative_to(root).as_posix(),
            'duration': duration, 'shots': records, 'fractions': fractions, 'narration_uses': 1,
            'timing': story['timing'] if cues else 'proportional_source_characters_not_word_aligned'})
    with tempfile.TemporaryDirectory(prefix='drama-ppt-', dir=root) as temp:
        raw = Path(temp)/'lesson.pptx'
        deck.save(raw)
        os.replace(raw, output)
    manifest = {'version': 4, 'template_id': lesson.ppt_template, 'style_mode': 'comic',
        'story_format': 'scene_drama', 'pages': pages, 'segments': segments, 'bookends': [],
        'media_fingerprint': media_fingerprint(lesson, audio_files), 'narration_uses_per_segment': 1,
        'preview_type': 'project_composition_not_powerpoint_screenshot',
        'motion': 'provider_video' if any(s.get('video') for items in lesson.deck_plan['generated_media'].values() for s in items)
                  else 'illustrated_still_sequence'}
    (previews/'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf8')
    return manifest


def drama_preview_svg():
    """A workflow diagram, explicitly not an example of generated anime art."""
    return '<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">' \
        '<rect width="1280" height="720" fill="#101623"/>' \
        '<g font-family="Microsoft YaHei,sans-serif" fill="#FFFFFF">' \
        '<text x="76" y="120" font-size="48" font-weight="bold">场景漫剧</text>' \
        '<text x="76" y="177" font-size="25" fill="#BBC8DC">原文旁白 · 动漫场景 · 连续分镜</text>' \
        '<path d="M78 330H1188" stroke="#7289B0" stroke-width="3"/>' \
        '<g font-size="30"><text x="78" y="300">远景建立环境</text>' \
        '<text x="468" y="300">中景表现行动</text><text x="858" y="300">特写突出细节</text></g>' \
        '<text x="76" y="452" font-size="28">人物、服装、地点随故事承接；按镜头绘制插画或生成视频</text>' \
        '<text x="76" y="534" font-size="25" fill="#BBC8DC">制作流程示意；接入绘图或视频服务后，逐镜生成实际动漫场景。</text>' \
        '<text x="76" y="647" font-size="20" fill="#BBC8DC">PPT收录场景关键帧，视频配完整教材原文与同步字幕</text></g></svg>'
