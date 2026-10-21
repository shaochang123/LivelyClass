"""Original vector characters, source-bound speech and shared comic media."""
from __future__ import annotations

import base64
import hashlib
import html
import json
import os
from pathlib import Path
import re
import tempfile
import wave

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from zhijiang.course_styles import ComicStyleDraft, effective_template, validate_panel_coverage, comic_sentences
from zhijiang.deck_renderer import Page, PAGE_W, PAGE_H
from zhijiang.presentation import PresentationError, _font, _notes, _patch_svg_parts
from zhijiang.presentation_templates import backdrop_elements
from zhijiang.page_media import narrated_pages, media_fingerprint, _run


def comic_lines(text, width, size=27, *, bold=False):
    """Measure the real font, preserving complete Latin words and all content."""
    font=_font(size,bold=bold);draw=ImageDraw.Draw(Image.new('RGB',(1,1)))
    tokens=re.findall(r"[A-Za-z0-9]+(?:['’_-][A-Za-z0-9]+)*|\s+|.",text)
    result=[];line=''
    for token in tokens:
        if token.isspace() and '\n' in token:
            if line:result.append(line.rstrip());line=''
            continue
        if draw.textlength(token,font=font)>width:
            raise PresentationError('漫剧对白的完整单词超出可读宽度，请选择更宽的分镜。')
        if line and draw.textlength(line+token,font=font)>width:
            result.append(line.rstrip());line=''
        if line or not token.isspace():line+=token
    if line:result.append(line.rstrip())
    return result or ['']


def _text(lines,x,y,*,size=27,color='172D35',bold=False,line_height=None):
    height=line_height or size*1.32
    return ''.join(f'<text x="{x}" y="{y+i*height}" font-family="Microsoft YaHei" '
        f'font-size="{size}" font-weight="{700 if bold else 400}" fill="#{color}">{html.escape(line)}</text>'
        for i,line in enumerate(lines))


def character_svg(character,x,y,scale=1,*,emotion='neutral',gesture='open',shot='medium',ink='172D35'):
    """Project-authored reusable cast; expressions/poses are declarative."""
    shirt=character['shirt'];hair=character['hair'];role=character['role']
    top='<g transform="translate('+str(x)+' '+str(y)+') scale('+str(scale)+')" stroke="#'+ink+'" stroke-linecap="round" stroke-linejoin="round">'
    parts=[top]
    if hair=='bob':parts.append(f'<path d="M-37-89Q-44-143 0-146Q44-143 37-89L32-67H-32Z" fill="#{ink}" stroke-width="3"/>')
    parts.append(f'<path d="M-44-48Q0-70 44-48L51 18H-51Z" fill="#{shirt}" stroke-width="4"/>')
    if shot=='wide':parts.append(f'<path d="M-27 20L-30 63M27 20L30 63M-42 65H-15M15 65H42" fill="none" stroke-width="12"/>')
    else:parts.append('<path d="M-49 20H49" fill="none" stroke-width="4"/>')
    parts.append('<ellipse cx="0" cy="-105" rx="34" ry="40" fill="#F4CCAD" stroke-width="4"/>')
    if hair=='short':parts.append(f'<path d="M-34-111Q-38-152 3-149Q41-148 34-116L18-133L-11-126L-24-134Z" fill="#{ink}" stroke-width="2"/>')
    elif hair=='curly':
        parts += [f'<circle cx="{cx}" cy="{cy}" r="13" fill="#{ink}" stroke-width="2"/>'
                  for cx,cy in [(-27,-135),(-10,-146),(12,-146),(29,-134)]]
    else:parts.append(f'<path d="M-34-110Q-37-150 0-147Q38-148 34-111L12-130L-12-124Z" fill="#{ink}" stroke-width="2"/>')
    parts += ['<circle cx="-12" cy="-105" r="3" fill="#172D35" stroke="none"/>',
              '<circle cx="12" cy="-105" r="3" fill="#172D35" stroke="none"/>']
    if role=='teacher':parts.append('<path d="M-25-114H-3V-98H-25ZM3-114H25V-98H3ZM-3-107H3" fill="none" stroke-width="2"/>')
    if emotion=='thinking':parts.append('<path d="M-17-120L-6-117M6-117L18-121M-7-85H7" fill="none" stroke-width="3"/>')
    elif emotion=='curious':parts.append('<path d="M-18-122L-7-120M7-123L18-126" fill="none" stroke-width="3"/><ellipse cx="1" cy="-86" rx="5" ry="6" fill="#FFF6E8" stroke-width="2"/>')
    else:parts.append('<path d="M-11-87Q0-76 12-87" fill="none" stroke-width="3"/>')
    arms={'point':'M-41-39L-67-16L-79-37M40-39L68-74L102-78',
          'think':'M-41-39L-58-4M40-39L59-62L27-86',
          'read':'M-41-39L-20-12M41-39L20-12',
          'open':'M-41-39L-73-13L-97-28M41-39L73-13L97-28'}
    parts.append(f'<path d="{arms[gesture]}" fill="none" stroke-width="9"/>')
    if gesture=='read':parts.append('<path d="M-34-25L0-15L34-25V17L0 27L-34 17Z" fill="#FFF9EC" stroke-width="3"/>')
    parts.append('</g>')
    return ''.join(parts)


def _embedded_svg(source,x,y,width,height):
    node=etree.fromstring(source)
    if node.get('viewBox') is None:node.set('viewBox',f"0 0 {node.get('width')} {node.get('height')}")
    for name,value in {'x':x,'y':y,'width':width,'height':height}.items():node.set(name,str(value))
    node.set('preserveAspectRatio','xMidYMid meet')
    return etree.tostring(node,encoding='unicode')


def _image_svg(path):
    with Image.open(path) as im:width,height=im.size
    data=base64.b64encode(Path(path).read_bytes()).decode()
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">'
        f'<image width="{width}" height="{height}" href="data:image/png;base64,{data}"/></svg>').encode()


def comic_svg(title,lines,panel,characters,template,board,*,page=1,source_page=None,wide=False,panorama=False):
    ink=template.ink;accent=template.accent
    cast={c['role']:c for c in characters};active=cast[panel['speaker']]
    title_lines=comic_lines(title,1150,32,bold=True)
    if len(title_lines)>3:raise PresentationError('漫剧标题超出三行，请拆分讲解主题。')
    top=max(128,40+len(title_lines)*40+24)
    if panorama:
        return panorama_svg(title_lines,lines,panel,cast,template,board,page,source_page)
    parts=['<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">',
        '<title>'+html.escape(title)+'</title><desc>原创矢量教学漫剧；对白来自对应讲稿，图示保留来源关系。</desc>',
        backdrop_elements(template),_text(title_lines,44,48,size=32,bold=True,color=ink),
        f'<rect x="36" y="{top}" width="760" height="{666-top}" rx="8" fill="#{template.panel}" stroke="#{ink}" stroke-width="4"/>',
        f'<rect x="816" y="{top+5}" width="428" height="{661-top}" rx="8" fill="#{template.background}" stroke="#{ink}" stroke-width="4"/>']
    if not wide:
        parts += [f'<rect x="165" y="{top+35}" width="605" height="345" rx="7" fill="#{template.background}" stroke="#{accent}" stroke-width="3"/>',
                  _embedded_svg(board,174,top+42,588,331),
                  character_svg(cast['teacher'],105,600,1,emotion='confident',gesture='point',shot='wide',ink=ink),
                  _text(comic_lines(cast['teacher']['name'],110,20,bold=True),45,430,size=20,bold=True,color=ink)]
        bx,by,bw,bh=843,top+29,374,303
    else:
        bx,by,bw,bh=70,top+30,1110,330
    parts += [f'<path d="M{bx+16} {by}H{bx+bw-16}Q{bx+bw} {by} {bx+bw} {by+16}V{by+bh-16}'
        f'Q{bx+bw} {by+bh} {bx+bw-16} {by+bh}H{bx+bw-94}L{bx+bw-51} {by+bh+35}L{bx+bw-132} {by+bh}'
        f'H{bx+16}Q{bx} {by+bh} {bx} {by+bh-16}V{by+16}Q{bx} {by} {bx+16} {by}Z" '
        f'fill="#{template.background}" stroke="#{ink}" stroke-width="3"/>',
        _text(lines,bx+20,by+38,size=27,color=ink,line_height=35)]
    scale={'wide':.73,'medium':.94,'close':1.13}[panel['shot']]
    parts += [character_svg(active,1110,614,scale,emotion=panel['emotion'],gesture=panel['gesture'],shot=panel['shot'],ink=ink),
              _text([active['name']],855,638,size=23,bold=True,color=ink)]
    source=f'教材第 {source_page} 页' if source_page is not None else 'Lively Class——AI教育Agent'
    parts += [_text([source],42,699,size=17,color=template.muted),
              _text([f'{page:02d}'],1190,699,size=17,color=template.muted),'</svg>']
    return ''.join(parts).encode()


def panorama_svg(title_lines,lines,panel,cast,template,board,page,source_page):
    """A dense board gets full width; speech continues in readable short shots."""
    ink=template.ink;active=cast[panel['speaker']]
    title_height=len(title_lines)*40
    top=max(116,40+title_height+12)
    board_height=min(390,536-top-12)
    if board_height<330:raise PresentationError('漫剧宽幅分镜需要更短的主题标题。')
    tail=('M172 597L137 592L172 616' if panel['speaker']=='teacher' else 'M1108 597L1150 592L1108 616')
    source=f'教材第 {source_page} 页' if source_page is not None else 'Lively Class——AI教育Agent'
    return ('<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">'
        '<title>'+html.escape(' '.join(title_lines))+'</title><desc>宽幅教学板书与角色对白分镜，图示内容完整保留。</desc>'
        +backdrop_elements(template)+_text(title_lines,44,48,size=32,bold=True,color=ink)
        +f'<rect x="36" y="{top}" width="1208" height="{board_height+24}" rx="8" fill="#{template.panel}" stroke="#{ink}" stroke-width="4"/>'
        +_embedded_svg(board,80,top+12,1120,board_height)
        +f'<rect x="36" y="548" width="1208" height="122" rx="8" fill="#{template.panel}" stroke="#{ink}" stroke-width="4"/>'
        +character_svg(cast['teacher'],92,648,.60,emotion=panel['emotion'] if panel['speaker']=='teacher' else 'thinking',gesture=panel['gesture'] if panel['speaker']=='teacher' else 'read',ink=ink)
        +character_svg(cast['student'],1185,648,.60,emotion=panel['emotion'] if panel['speaker']=='student' else 'curious',gesture=panel['gesture'] if panel['speaker']=='student' else 'think',ink=ink)
        +f'<rect x="172" y="553" width="936" height="110" rx="12" fill="#{template.background}" stroke="#{ink}" stroke-width="3"/>'
        +f'<path d="{tail}" fill="#{template.background}" stroke="#{ink}" stroke-width="3"/>'
        +_text([active['name']],190,578,size=20,bold=True,color=ink)
        +_text(lines,190,610,size=27,color=ink,line_height=34)
        +_text([source],42,699,size=17,color=template.muted)
        +_text([f'{page:02d}'],1190,699,size=17,color=template.muted)+'</svg>').encode()


def comic_preview_svg(template):
    cast=[dict(role='teacher',name='讲师',hair='short',shirt=template.accent),
          dict(role='student',name='学员',hair='bob',shirt=template.secondary)]
    board=(f'<svg xmlns="http://www.w3.org/2000/svg" width="588" height="331" viewBox="0 0 588 331">'
        f'<path d="M80 210H500M100 250V50" stroke="#{template.muted}" stroke-width="3"/>'
        f'<path d="M100 210L200 160L300 160L400 80L500 120" stroke="#{template.accent}" stroke-width="6" fill="none"/>'
        +_text(['教学图示保留在场景中'],90,300,size=25,color=template.ink)+'</svg>').encode()
    return comic_svg('漫剧讲解 角色与分镜示意',['角色讲解教材内容，','对白随讲稿逐镜呈现。'],
        dict(speaker='student',shot='medium',emotion='curious',gesture='think'),cast,template,board).decode()


def bookend_svg(title,lines,characters,template,page):
    title_lines=comic_lines(title,1090,40,bold=True)
    if len(title_lines)>3:raise PresentationError('漫剧封面标题过长，请拆分主题。')
    body_y=max(280,100+len(title_lines)*53+24)
    return ('<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">'
        '<title>'+html.escape(title)+'</title>'+backdrop_elements(template)
        +f'<rect x="36" y="64" width="1208" height="600" rx="10" fill="#{template.panel}" stroke="#{template.ink}" stroke-width="4"/>'
        +_text(title_lines,80,123,size=40,bold=True,color=template.ink,line_height=53)
        +_text(lines,80,body_y,size=27,color=template.ink,line_height=35)
        +character_svg(characters[0],1040,455,1.06,emotion='confident',gesture='open',shot='wide',ink=template.ink)
        +character_svg(characters[1],1110,615,.82,emotion='curious',gesture='think',shot='medium',ink=template.ink)
        +_text(['Lively Class——AI教育Agent'],42,699,size=17,color=template.muted)
        +_text([f'{page:02d}'],1190,699,size=17,color=template.muted)+'</svg>').encode()


def _knowledge_boards(lesson,index,template,root,media):
    segment=lesson.segments[index-1]
    computed=bool(segment.math_scene or segment.visual_scene or index in lesson.animation_report.get('prepared_scenes',[]))
    if computed:
        clip=root/('math' if segment.math_scene else 'visual')/f'scene-{index:02d}'/'clip.mp4'
        if not clip.is_file():raise PresentationError('漫剧引用的连续演示尚未生成。')
        poster=media/f'computed-{index:03d}.png'
        _run(['-sseof','-0.08','-i',str(clip),'-frames:v','1','-update','1',str(poster)],poster)
        return [_image_svg(poster)],'computed_scene',clip
    visual=lesson.deck_plan.get('visual_story',{}).get('visuals',{}).get(str(index),{}).get('visual')
    if visual:
        from zhijiang.story_svg import visual_pages, visual_svg
        try:
            items=visual_pages(visual,template,588,331)
            return [visual_svg(v,template,588,331)[0] for v in items],'teaching_svg',None
        except PresentationError:
            # Width is a drawing constraint, not a reason to omit a condition.
            items=visual_pages(visual,template,1120,390)
            return [visual_svg(v,template,1120,390)[0] for v in items],'teaching_svg_panorama',None
    figure=lesson.deck_plan.get('source_figures',{}).get(str(index))
    if figure and (root/figure['file']).is_file():return [_image_svg(root/figure['file'])],'source_figure',None
    # Existing bullets are retained verbatim; this is an explanatory board,
    # never a fabricated subject diagram.
    lines=[]
    for bullet in segment.bullets:lines.extend(comic_lines(bullet,548,28))
    boards=[]
    for start in range(0,len(lines),7):
        boards.append(('<svg xmlns="http://www.w3.org/2000/svg" width="588" height="331" viewBox="0 0 588 331">'
            +_text(lines[start:start+7],20,45,size=28,color=template.ink,line_height=40)+'</svg>').encode())
    return boards,'existing_bullets',None


def render_comic_deck(lesson,output,*,audio_files=None):
    if lesson.deck_plan.get('comic_story',{}).get('format')=='scene_drama':
        from zhijiang.drama_renderer import render_drama_deck
        return render_drama_deck(lesson,output,audio_files=audio_files)
    import resvg_py
    from zhijiang.page_media import scene_on_video
    output=Path(output);root=output.parent;root.mkdir(parents=True,exist_ok=True)
    previews=root/'presentation-preview';media=root/'presentation-media'
    previews.mkdir(exist_ok=True);media.mkdir(exist_ok=True)
    template=effective_template(lesson)
    style=ComicStyleDraft.model_validate(lesson.deck_plan['course_style']['design'])
    cast=[c.model_dump() for c in style.characters]
    story=lesson.deck_plan['comic_story']
    if story['characters']!=cast:raise PresentationError('漫剧角色与课程风格记录不一致。')
    if audio_files is None:
        audio_files=[root/('math' if s.math_scene else 'visual')/f'scene-{i:02d}'/'narration.wav'
            if s.math_scene or s.visual_scene else root/f'audio-{i-1}.wav' for i,s in enumerate(lesson.segments,1)]
    if len(audio_files)!=len(lesson.segments):raise PresentationError('漫剧配音数量与课程不一致。')
    deck=Presentation();deck.slide_width=Inches(PAGE_W);deck.slide_height=Inches(PAGE_H)
    deck.core_properties.title=lesson.title;deck.core_properties.subject='Lively Class——AI教育Agent · 原创矢量漫剧'
    assets={};pages=[];segments=[];bookends=[]

    def add_page(svg,kind,index=None,panel=None,segment=None):
        num=len(deck.slides)+1
        svgpath=media/f'comic-{num:03d}.svg';png=previews/f'page-{num:03d}.png'
        svgpath.write_bytes(svg);png.write_bytes(resvg_py.svg_to_bytes(svg_string=svg.decode(),font_family='Microsoft YaHei'))
        page=Page(deck,template,previews,num);shape=page.picture(png,(0,0,PAGE_W,PAGE_H))
        assets[num]=[(shape.shape_id,svg)]
        if segment is not None:_notes(page.slide,segment,animation=False)
        page.slide.notes_slide.notes_text_frame.text+='\n漫剧分镜：'+json.dumps(panel or {},ensure_ascii=False)
        pages.append({'page':num,'kind':kind,'segment_id':index,'panel':panel,
                      'svg':svgpath.relative_to(root).as_posix()})
        return png,num

    bookend_pages=[]
    objective_lines=comic_lines(lesson.objective,760)
    for start in range(0,len(objective_lines),8):
        cover,_=add_page(bookend_svg(lesson.title,objective_lines[start:start+8],cast,template,len(deck.slides)+1),'cover')
        deck.slides[-1].notes_slide.notes_text_frame.text=lesson.objective+'\n'+lesson.notice
        bookend_pages.append(('cover',cover))
    for index,(segment,audio) in enumerate(zip(lesson.segments,audio_files),1):
        entry=story['segments'].get(str(index),{})
        if entry.get('narration_sha256')!=hashlib.sha256(segment.narration.encode()).hexdigest():
            raise PresentationError('讲稿已经变化，需要重新规划漫剧分镜。')
        panels=entry['panels'];sentences=comic_sentences(segment.narration)
        validate_panel_coverage(panels,range(1,len(sentences)+1))
        for panel in panels:
            if panel['text']!=''.join(sentences[n-1] for n in panel['sentence_ids']):
                raise PresentationError('漫剧对白与原讲稿不一致。')
        boards,board_source,source_clip=_knowledge_boards(lesson,index,template,root,media)
        panorama=board_source=='teaching_svg_panorama'
        if not panorama:
            try:
                for panel in panels:comic_lines(panel['text'],334)
            except PresentationError:
                panorama=True
        states=[];weights=[];numbers=[];panel_records=[]
        for pnum,panel in enumerate(panels):
            wide=False
            lines=comic_lines(panel['text'],896 if panorama else 334)
            capacity=2 if panorama else 7
            chunks=[lines[n:n+capacity] for n in range(0,len(lines),capacity)]
            count=max(len(chunks),len(boards) if pnum==0 else 1)
            for part in range(count):
                text=chunks[min(part,len(chunks)-1)]
                board=boards[part%len(boards)]
                png,number=add_page(comic_svg(segment.title,text,panel,cast,template,board,
                    page=len(deck.slides)+1,source_page=segment.evidence.page,wide=wide,panorama=panorama),
                    'comic',index,panel,segment)
                states.append(png);numbers.append(number)
                weights.append(max(1,len(panel['text']))/count)
                panel_records.append({'page':number,'speaker':panel['speaker'],'sentence_ids':panel['sentence_ids'],
                    'shot':panel['shot'],'emotion':panel['emotion'],'gesture':panel['gesture'],
                    'board_source':board_source,'text_fragment':'\n'.join(text)})
        generated_windows=[]
        for item in lesson.deck_plan.get('generated_media',{}).get(str(index),[]):
            poster=(root/item['file']).resolve()
            if not poster.is_relative_to(root.resolve()) or not poster.is_file():raise PresentationError('生成插画路径无效。')
            art=Page(deck,template,previews,len(deck.slides)+1)
            art.text(segment.title,(.6,.25,12.1,.8),size=30,min_size=30,bold=True)
            box=(.6,1.2,12.1,5.7);art.picture(poster,box)
            art.text('AI原创教学插画 · 非教材原图',(.6,7.,12.,.3),size=12,min_size=12)
            _notes(art.slide,segment,animation=False);art.save_preview()
            pages.append({'page':art.number,'kind':'generated_illustration','segment_id':index})
            states.append(previews/f'page-{art.number:03d}.png');numbers.append(art.number);weights.append(max(1,len(segment.narration))/max(1,len(states)))
            generated_windows.append({'state':len(states)-1,'box':[v*96 for v in box],**item})
        cumulative=0;total=sum(weights);fractions=[]
        for weight in weights:fractions.append(cumulative/total);cumulative+=weight
        clip=media/f'segment-{index:03d}.mp4'
        if not Path(audio).is_file():raise PresentationError('漫剧缺少对应讲稿的真实配音。')
        if source_clip:
            background=media/f'comic-background-{index:03d}.mp4'
            duration=narrated_pages(states,audio,background,fractions=fractions)
            title_lines=comic_lines(segment.title,1150,32,bold=True)
            top=max(128,40+len(title_lines)*40+24)
            scene_on_video(background,source_clip,(174,top+42,588,331),clip,duration)
        else:duration=narrated_pages(states,audio,clip,fractions=fractions)
        for n,item in enumerate(generated_windows):
            if not source_clip and not item.get('video'):continue
            from zhijiang.page_media import overlay_generated_window
            target=media/f'art-overlay-{index}-{n}.mp4';source=root/item.get('video',item['file'])
            if not source.resolve().is_relative_to(root.resolve()):raise PresentationError('媒体路径越界。')
            start=fractions[item['state']];end=fractions[item['state']+1] if item['state']+1<len(fractions) else 1
            overlay_generated_window(clip,source,item['box'],target,start*duration,end*duration,duration,
                still=not bool(item.get('video')));os.replace(target,clip)
        segments.append({'segment_id':index,'pages':numbers,'clip':clip.relative_to(root).as_posix(),
            'duration':duration,'panels':panel_records,'timing':story['timing'],'fractions':fractions,
            'narration_uses':1,'continuous_scene_retained':source_clip is not None,'generated_media':generated_windows})
    last=lesson.segments[-1]
    closing_lines=comic_lines(last.bullets[-1],760)
    for start in range(0,len(closing_lines),8):
        ending,_=add_page(bookend_svg('本课回顾',closing_lines[start:start+8],cast,template,len(deck.slides)+1),'ending')
        _notes(deck.slides[-1],last,animation=False)
        bookend_pages.append(('ending',ending))
    for number,(kind,png) in enumerate(bookend_pages):
        silence=media/f'comic-{kind}-{number}.wav'
        with wave.open(str(silence),'wb') as w:
            w.setnchannels(1);w.setsampwidth(2);w.setframerate(24000);w.writeframes(b'\x00\x00'*24000)
        clip=media/f'comic-{kind}-{number}.mp4';narrated_pages([png],silence,clip)
        bookends.append({'kind':kind,'clip':clip.relative_to(root).as_posix(),'duration':1,
            'audio':'intentional_silence_no_repeated_narration'})
    with tempfile.TemporaryDirectory(prefix='comic-ppt-',dir=root) as temp:
        raw=Path(temp)/'raw.pptx';final=Path(temp)/'final.pptx'
        deck.save(raw);_patch_svg_parts(raw,final,assets);os.replace(final,output)
    manifest={'version':3,'template_id':lesson.ppt_template,'style_mode':'comic','pages':pages,
        'segments':segments,'bookends':bookends,'media_fingerprint':media_fingerprint(lesson,audio_files),
        'preview_type':'project-rendered-svg-and-png-not-powerpoint-screenshot',
        'cast':cast,'narration_uses_per_segment':1}
    (previews/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return manifest
